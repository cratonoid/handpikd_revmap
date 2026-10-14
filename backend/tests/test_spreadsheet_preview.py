# The in-browser Excel viewer (services/spreadsheet_preview.py): cells must
# read the way Excel shows them, hidden sheets stay hidden, and the empty
# margins Excel likes to keep around a sheet are trimmed off.
import io
from datetime import datetime

from openpyxl import Workbook

from app.services import spreadsheet_preview
from app.services.spreadsheet_preview import format_cell, read_workbook_preview


def _xlsx(build) -> bytes:
    workbook = Workbook()
    build(workbook)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_cells_read_the_way_excel_shows_them():
    assert format_cell(None) == ""
    assert format_cell(120.0) == "120"
    assert format_cell(0.1 + 0.2) == "0.3"
    assert format_cell(1234.5) == "1234.5"
    assert format_cell(0.18, "0%") == "18%"
    assert format_cell(datetime(2026, 4, 1)) == "01/04/2026"
    assert format_cell(True) == "TRUE"
    assert format_cell("  Mug 350ml ") == "Mug 350ml"


def test_xlsx_sheets_rows_and_values():
    def build(workbook):
        sheet = workbook.active
        sheet.title = "Drinkware"
        sheet.append(["Item", "MRP", "Rate"])
        sheet.append(["Mug", 250, 180.5])
        sheet["B3"] = "=B2*2"  # never calculated by Excel, so no cached value
        workbook.create_sheet("Bags").append(["Tote", 400])
        workbook.create_sheet("Internal").sheet_state = "hidden"

    sheets = read_workbook_preview(_xlsx(build), "xlsx")

    assert [sheet.name for sheet in sheets] == ["Drinkware", "Bags"]
    assert sheets[0].rows[:2] == [["Item", "MRP", "Rate"], ["Mug", "250", "180.5"]]
    assert sheets[1].rows == [["Tote", "400"]]
    assert not sheets[0].truncated


def test_trailing_empty_rows_and_columns_are_trimmed():
    def build(workbook):
        sheet = workbook.active
        sheet["A1"] = "Item"
        sheet["B2"] = 10
        sheet["F20"] = None
        sheet["F20"].number_format = "0.00"  # formatted but empty: Excel widens the sheet for it

    sheets = read_workbook_preview(_xlsx(build), "xlsx")
    assert sheets[0].rows == [["Item", ""], ["", "10"]]


def test_large_sheets_are_cut_off_and_flagged(monkeypatch):
    monkeypatch.setattr(spreadsheet_preview, "MAX_ROWS", 3)

    def build(workbook):
        for n in range(5):
            workbook.active.append([f"row {n}"])

    sheet = read_workbook_preview(_xlsx(build), "xlsx")[0]
    assert sheet.rows == [["row 0"], ["row 1"], ["row 2"]]
    assert sheet.truncated
